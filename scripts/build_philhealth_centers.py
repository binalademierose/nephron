"""
build_philhealth_centers.py - turns the PhilHealth-provided facility list (an .xlsx) into
data/philhealth_centers.json, the file services/philhealth_centers.py actually reads at runtime.

Run once, whenever PhilHealth provides an updated list:
    pip install openpyxl --break-system-packages   # if not already installed
    python scripts/build_philhealth_centers.py path/to/the-new-list.xlsx

The source sheet is not a plain table: a "divider" row (only its first cell filled in) introduces either a
region (its text contains the word "REGION") or, until the next one, a province or an independent city; every
other row is one facility, in whichever region/province-or-city came before it. This script reads that structure
literally - it never invents a facility, an address, or an accreditation detail that is not in the source file.
"""
import json
import os
import re
import sys
import unicodedata

SECTOR_LABELS = {"P": "Private", "G": "Government"}


def _is_region(text):
    return "REGION" in text.upper()


def _clean(value):
    if value is None:
        return None
    return re.sub(r"\s+", " ", str(value)).strip()


def build(xlsx_path):
    import openpyxl

    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(min_row=2, values_only=True)   # skip the header row

    centers = []
    current_region, current_area, next_id = None, None, 1
    for row in rows:
        first = row[0]
        if first is None:
            continue
        if isinstance(first, str) and all(cell is None for cell in row[1:]):
            text = _clean(first)
            if _is_region(text):
                current_region, current_area = text, None
            else:
                current_area = text
            continue
        if isinstance(first, (int, float)):
            name, tel, email, street, municipality, expire_date, sector = row[1:8]
            centers.append({
                "id": next_id,
                "name": _clean(name),
                "region": current_region,
                "province_or_city": current_area or _clean(municipality),
                "municipality": _clean(municipality),
                "street": _clean(street),
                "tel": _clean(tel),
                "email": _clean(email) if email else None,
                "expire_date": expire_date.date().isoformat() if hasattr(expire_date, "date") else None,
                "sector": SECTOR_LABELS.get(sector, sector),
            })
            next_id += 1
    return centers


def main():
    if len(sys.argv) != 2:
        print("Usage: python scripts/build_philhealth_centers.py path/to/source.xlsx")
        raise SystemExit(1)
    centers = build(sys.argv[1])
    out_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "philhealth_centers.json")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as handle:
        json.dump({"source": f"PhilHealth accredited dialysis/kidney facility directory ({os.path.basename(sys.argv[1])})",
                  "centers": centers}, handle, indent=2, ensure_ascii=False)
    print(f"Wrote {len(centers)} facilities to {out_path}")


if __name__ == "__main__":
    main()
